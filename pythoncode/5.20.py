import pandas as pd
import numpy as np
from sklearn.model_selection import train_test_split
from sklearn.ensemble import RandomForestRegressor
from sklearn.multioutput import MultiOutputRegressor
from sklearn.preprocessing import StandardScaler
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import logging
import os
plt.rcParams['font.sans-serif'] = ['SimHei']
plt.rcParams['axes.unicode_minus'] = False

# 配置日志
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

# 全局变量
prediction_results = pd.DataFrame()
metrics_results = {}

# =================== 配置参数 ===================
DATA_PATH = r'C:\Users\admin\Desktop\NSW半小时时电价预测数据管理 (1).csv'
PREDICTION_DATE = '2024-04-11'
OUTPUT_EXCEL = '预测结果_评估指标.xlsx'
OUTPUT_IMAGE = '预测结果可视化.png'

LAGS = [2, 4]
NUMERIC_COLS = ['Temperature', 'WindSpeed', 'Cloud', 'Humidity', 'Rain']
TARGET_FEATURES = ['SpotPrice_RRP', 'TotalDemand']


# =================== 工具函数 ===================
def calculate_metrics(y_true, y_pred, feature_name):
    metrics = {}
    metrics['MAE'] = mean_absolute_error(y_true, y_pred)
    metrics['RMSE'] = np.sqrt(mean_squared_error(y_true, y_pred))
    metrics['R2'] = r2_score(y_true, y_pred)

    value_range = np.max(y_true) - np.min(y_true)
    metrics['NMAE'] = metrics['MAE'] / value_range if value_range != 0 else np.nan
    metrics['NRMSE'] = metrics['RMSE'] / value_range if value_range != 0 else np.nan

    epsilon = 1e-10
    with np.errstate(divide='ignore', invalid='ignore'):
        mape = np.nanmean(np.abs((y_true - y_pred) / (y_true + epsilon))) * 100
    metrics['MAPE(%)'] = mape

    imape_values = []
    for true_val, pred_val in zip(y_true, y_pred):
        if "Price" in feature_name:
            if true_val < 0:
                if pred_val < 0:
                    imape_values.append(0)
                elif 0 <= pred_val <= 30:
                    imape_values.append(0.2)
                else:
                    imape_values.append(1)
            else:
                imape_values.append(min(np.abs((true_val - pred_val) / (true_val + epsilon)), 1))
        else:
            imape_values.append(min(np.abs((true_val - pred_val) / (true_val + epsilon)), 1))

    metrics['IMAPE'] = np.nanmean(imape_values)
    return metrics


def safe_create_chart(sheet, data_range, title, y_title, position, style=13):
    try:
        chart = LineChart()
        chart.add_data(data_range, titles_from_data=True)
        chart.title = title
        chart.style = style
        chart.y_axis.title = y_title

        date_axis = DateAxis(crossAx=100)
        date_axis.number_format = 'yyyy-mm-dd hh:mm'
        date_axis.title = "时间"
        chart.x_axis = date_axis

        sheet.add_chart(chart, position)
        return True
    except Exception as e:
        logging.error(f"图表创建失败: {str(e)}")
        return False


# =================== 数据加载与预处理 ===================
def load_and_preprocess_data(file_path):
    df = pd.read_csv(file_path, skiprows=1)
    df['DateTime'] = pd.to_datetime(df['DateTime'], format='%Y/%m/%d %H:%M')
    df.set_index('DateTime', inplace=True)
    df['WeatherType'] = df['WeatherType'].ffill()
    return df


def build_features(df):
    df['hour'] = df.index.hour
    df['minute'] = df.index.minute
    df['dayofweek'] = df.index.dayofweek
    df['is_weekend'] = df['dayofweek'].apply(lambda x: 1 if x >= 5 else 0)

    weather_dummies = pd.get_dummies(df['WeatherType'].astype(str), prefix='weather')
    df = pd.concat([df, weather_dummies], axis=1)

    for lag in LAGS:
        df[f'RRP_lag{lag}'] = df['SpotPrice_RRP'].shift(lag)
        df[f'Demand_lag{lag}'] = df['TotalDemand'].shift(lag)
    df.dropna(inplace=True)

    for col in NUMERIC_COLS:
        df[col] = pd.to_numeric(df[col], errors='coerce').ffill()

    features = ['hour', 'minute', 'dayofweek', 'is_weekend', 'Temperature',
                'WindSpeed', 'Cloud', 'Humidity', 'Rain'] + \
               [col for col in df.columns if col.startswith('weather_')] + \
               [f'RRP_lag{lag}' for lag in LAGS] + \
               [f'Demand_lag{lag}' for lag in LAGS]

    X = df[features].values
    y = df[TARGET_FEATURES].values

    return df, X, y, features


# =================== 模型训练与评估 ===================
def train_model(X_train, y_train):
    model = MultiOutputRegressor(
        RandomForestRegressor(n_estimators=200, max_depth=10, random_state=42, n_jobs=-1)
    )
    model.fit(X_train, y_train)
    return model


def evaluate_model(model, X_test, y_test, scaler_y):
    y_pred_scaled = model.predict(X_test)
    y_pred = scaler_y.inverse_transform(y_pred_scaled)
    y_true = scaler_y.inverse_transform(y_test)

    for idx, feature in enumerate(TARGET_FEATURES):
        metrics_results[feature] = calculate_metrics(y_true[:, idx], y_pred[:, idx], feature)


# =================== 预测未来数据 ===================
def predict_future(df, features, model, scaler_X, scaler_y):
    future_dates = pd.date_range(start=f'{PREDICTION_DATE} 00:00',
                                 end=f'{PREDICTION_DATE} 23:30',
                                 freq='30T')
    future_df = pd.DataFrame(index=future_dates)

    future_df['hour'] = future_df.index.hour
    future_df['minute'] = future_df.index.minute
    future_df['dayofweek'] = future_df.index.dayofweek
    future_df['is_weekend'] = future_df['dayofweek'].apply(lambda x: 1 if x >= 5 else 0)

    future_df['SpotPrice_RRP'] = np.nan
    future_df['TotalDemand'] = np.nan

    last_date = df.index[-1]
    seven_days_ago = last_date - pd.DateOffset(days=7)
    historical_mean = df.loc[seven_days_ago:last_date].mean(numeric_only=True)

    for col in NUMERIC_COLS:
        future_df[col] = historical_mean[col]

    weather_cols = [col for col in df.columns if col.startswith('weather_')]
    future_df[weather_cols] = 0
    if df['WeatherType'].notna().any():
        last_weather = df['WeatherType'].dropna().mode()[0]
        future_df[f'weather_{last_weather}'] = 1

    last_4_steps = df[TARGET_FEATURES].iloc[-4:].values

    for i in range(len(future_df)):
        if i < 2:
            future_df.loc[future_df.index[i], 'RRP_lag2'] = last_4_steps[-2 + i, 0]
            future_df.loc[future_df.index[i], 'Demand_lag2'] = last_4_steps[-2 + i, 1]
            future_df.loc[future_df.index[i], 'RRP_lag4'] = last_4_steps[-4 + i, 0]
            future_df.loc[future_df.index[i], 'Demand_lag4'] = last_4_steps[-4 + i, 1]
        else:
            future_df.loc[future_df.index[i], 'RRP_lag2'] = future_df.loc[future_df.index[i - 2], 'SpotPrice_RRP']
            future_df.loc[future_df.index[i], 'Demand_lag2'] = future_df.loc[future_df.index[i - 2], 'TotalDemand']
            future_df.loc[future_df.index[i], 'RRP_lag4'] = future_df.loc[future_df.index[i - 4], 'SpotPrice_RRP']
            future_df.loc[future_df.index[i], 'Demand_lag4'] = future_df.loc[future_df.index[i - 4], 'TotalDemand']

        if i >= 2:
            current_features = future_df.loc[[future_df.index[i]], features].values
            current_pred = model.predict(scaler_X.transform(current_features))
            unscaled_pred = scaler_y.inverse_transform(current_pred)
            future_df.loc[future_df.index[i], 'SpotPrice_RRP'] = unscaled_pred[0, 0]
            future_df.loc[future_df.index[i], 'TotalDemand'] = unscaled_pred[0, 1]

    # 预测完成后填充剩余的 NaN
    future_df.ffill(inplace=True)  # 前向填充
    future_df.fillna(historical_mean, inplace=True)  # 再填充一次平均值

    global prediction_results
    prediction_results = pd.DataFrame({
        'DateTime': future_dates,
        'SpotPrice_RRP': future_df['SpotPrice_RRP'],
        'TotalDemand': future_df['TotalDemand']
    })


# =================== 结果保存 ===================
def save_results_to_excel(prediction_df, metrics_dict, file_path):
    try:
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            workbook = writer.book
            if not workbook.sheetnames:
                workbook.create_sheet("预测结果")

            prediction_df.to_excel(writer, sheet_name='预测结果', index=False)
            sheet = writer.sheets['预测结果']

            # 创建电价图表
            price_data = Reference(sheet, min_col=2, min_row=1, max_row=len(prediction_df)+1)
            date_data = Reference(sheet, min_col=1, min_row=2, max_row=len(prediction_df)+1)

            price_chart = LineChart()
            price_chart.title = "电价预测趋势"
            price_chart.y_axis.title = "价格 (AUD/MWh)"
            price_chart.add_data(price_data, titles_from_data=True)
            price_chart.set_categories(date_data)

            # 使用默认日期轴格式
            price_chart.x_axis.number_format = 'yyyy-mm-dd hh:mm'
            price_chart.x_axis.majorTimeUnit = "days"

            sheet.add_chart(price_chart, "E2")


            # 创建负荷图表
            load_data = Reference(sheet, min_col=3, min_row=1, max_row=len(prediction_df)+1)

            load_chart = LineChart()
            load_chart.title = "负荷预测趋势"
            load_chart.y_axis.title = "负荷 (MW)"
            load_chart.add_data(load_data, titles_from_data=True)
            load_chart.set_categories(date_data)

            sheet.add_chart(load_chart, "E20")

            # 保存评估指标
            for feature in metrics_dict:
                temp_df = pd.DataFrame.from_dict(metrics_dict[feature], orient='index', columns=['值'])
                temp_df.index.name = '指标'
                temp_df.to_excel(writer, sheet_name=f'{feature}指标')

            sheet.column_dimensions['A'].width = 20
            sheet.column_dimensions['B'].width = 15
            sheet.column_dimensions['C'].width = 15

        logging.info(f"结果已成功保存至 {file_path}")
    except Exception as e:
        logging.error(f"文件保存失败: {str(e)}")


# =================== 可视化 ===================
def plot_predictions(prediction_df):
    if prediction_df.empty:
        logging.warning("未生成预测结果，无法绘图")
        return

    plt.figure(figsize=(18, 12))
    plt.subplot(2, 1, 1)
    plt.plot(prediction_df['DateTime'], prediction_df['SpotPrice_RRP'])
    plt.title('电价预测')
    plt.gca().xaxis.set_major_formatter(mdates.DateFormatter('%H:%M'))

    plt.subplot(2, 1, 2)
    plt.plot(prediction_df['DateTime'], prediction_df['TotalDemand'], color='orange')
    plt.title('负荷预测')
    plt.gca().xaxis.set_major_formatter(mdates.DateFormatter('%H:%M'))

    plt.tight_layout()
    plt.savefig(OUTPUT_IMAGE, dpi=300, bbox_inches='tight')
    plt.close()
    logging.info(f"图表已保存至 {OUTPUT_IMAGE}")


# =================== 主流程 ===================
def main():
    try:
        df = load_and_preprocess_data(DATA_PATH)
        df, X, y, features = build_features(df)

        scaler_X = StandardScaler()
        scaler_y = StandardScaler()
        X_scaled = scaler_X.fit_transform(X)
        y_scaled = scaler_y.fit_transform(y)

        X_train, X_test, y_train, y_test = train_test_split(
            X_scaled, y_scaled, test_size=48*2, shuffle=False
        )

        model = train_model(X_train, y_train)
        evaluate_model(model, X_test, y_test, scaler_y)
        predict_future(df, features, model, scaler_X, scaler_y)
        save_results_to_excel(prediction_results, metrics_results, OUTPUT_EXCEL)
        plot_predictions(prediction_results)

    except Exception as e:
        logging.error(f"主程序发生错误: {str(e)}")
        with pd.ExcelWriter(OUTPUT_EXCEL) as writer:
            pd.DataFrame({'错误信息': [str(e)]}).to_excel(writer, sheet_name='错误日志')


if __name__ == '__main__':
    main()